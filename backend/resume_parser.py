import pdfplumber
import docx
import re
from skills_data import KNOWN_SKILLS, SKILL_ALIASES

def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

def extract_text_from_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    text_parts = []

    # regular paragraphs (outside tables)
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)

    # table content — many resume templates put sections inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)

    return "\n".join(text_parts)

def extract_text(file_path: str) -> str:
    if file_path.lower().endswith(".pdf"):
        return extract_text_from_pdf(file_path)
    elif file_path.lower().endswith(".docx"):
        return extract_text_from_docx(file_path)
    else:
        raise ValueError("Unsupported file type")

def extract_skills(text: str) -> list[str]:
    text_lower = text.lower()
    found_skills = set()  # use a set to automatically avoid duplicates

    # 1. Check direct matches against the main skill list
    for skill in KNOWN_SKILLS:
        pattern = r"\b" + re.escape(skill.lower()) + r"\b"
        if re.search(pattern, text_lower):
            found_skills.add(skill)

    # 2. Check aliases (e.g. "ml" in resume -> "Machine Learning")
    for alias, canonical_skill in SKILL_ALIASES.items():
        pattern = r"\b" + re.escape(alias) + r"\b"
        if re.search(pattern, text_lower):
            found_skills.add(canonical_skill)

    return sorted(found_skills)

def extract_email(text: str) -> str | None:
    match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", text)
    return match.group(0) if match else None

def extract_phone(text: str) -> str | None:
    match = re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}", text)
    return match.group(0) if match else None